#!/usr/bin/env python3
"""Build the AMKAD AR Reporting Transformation Project — internship close-out deck."""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------- palette ----------
RED = RGBColor(0xD4, 0x05, 0x11)
RED_DARK = RGBColor(0xA5, 0x04, 0x0D)
YELLOW = RGBColor(0xFF, 0xCC, 0x00)
INK = RGBColor(0x1A, 0x1D, 0x21)
SOFT = RGBColor(0x5B, 0x64, 0x70)
MUTED = RGBColor(0x8A, 0x83, 0x7B)
BG = RGBColor(0xF7, 0xF6, 0xF3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xE2, 0xE0, 0xDA)
GOOD = RGBColor(0x0F, 0x8A, 0x2E)
CRIT = RGBColor(0xCF, 0x2E, 0x2E)

FONT = "Calibri"
SCRATCH = "/tmp/claude-0/-home-user-Power-BI-/94de1877-f729-5659-9e51-b28f0bb44218/scratchpad/"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height


def add_slide(bg=BG):
    s = prs.slides.add_slide(BLANK)
    rect = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    rect.fill.solid(); rect.fill.fore_color.rgb = bg
    rect.line.fill.background()
    rect.shadow.inherit = False
    s.shapes._spTree.remove(rect._element)
    s.shapes._spTree.insert(2, rect._element)
    return s


def rect(s, l, t, w, h, color, line=False):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    r.fill.solid(); r.fill.fore_color.rgb = color
    if line:
        r.line.color.rgb = LINE; r.line.width = Pt(0.75)
    else:
        r.line.fill.background()
    r.shadow.inherit = False
    return r


def rrect(s, l, t, w, h, color, line=False, radius=0.06):
    r = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    try:
        r.adjustments[0] = radius
    except Exception:
        pass
    r.fill.solid(); r.fill.fore_color.rgb = color
    if line:
        r.line.color.rgb = LINE; r.line.width = Pt(0.75)
    else:
        r.line.fill.background()
    r.shadow.inherit = False
    return r


def textbox(s, l, t, w, h, valign=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    return tb, tf


def para(tf, text, size=14, color=INK, bold=False, italic=False, align=PP_ALIGN.LEFT,
         space_after=6, first=False, bullet=False, font=FONT):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    if bullet:
        text = "•  " + text
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r.font.color.rgb = color; r.font.name = font
    return p


def eyebrow(s, text, top=Inches(0.5), left=Inches(0.7)):
    tb, tf = textbox(s, left, top, Inches(8), Inches(0.4))
    para(tf, text.upper(), size=12.5, color=RED, bold=True, first=True)
    bar = rect(s, left, top + Inches(0.35), Inches(0.4), Pt(3.2), YELLOW)
    return bar


def slide_title(s, kicker, title, subtitle=None):
    eyebrow(s, kicker)
    tb, tf = textbox(s, Inches(0.7), Inches(0.9), Inches(11.9), Inches(1.1))
    para(tf, title, size=30, color=INK, bold=True, first=True)
    if subtitle:
        tb2, tf2 = textbox(s, Inches(0.7), Inches(1.55), Inches(11.9), Inches(0.6))
        para(tf2, subtitle, size=14.5, color=SOFT, first=True)
    rect(s, Inches(0.7), Inches(2.05), Inches(11.9), Pt(1.4), LINE)


def footer(s, n):
    tb, tf = textbox(s, Inches(0.7), Inches(7.15), Inches(6), Inches(0.3))
    para(tf, "AMKAD AR Reporting Transformation  ·  Internship Close-Out", size=9, color=MUTED, first=True)
    tb2, tf2 = textbox(s, Inches(12.4), Inches(7.15), Inches(0.6), Inches(0.3))
    para(tf2, str(n), size=9, color=MUTED, first=True, align=PP_ALIGN.RIGHT)


def bullets(s, items, left=Inches(0.7), top=Inches(2.35), width=Inches(11.9), size=15,
            gap=10, sub_size=12.5):
    tb, tf = textbox(s, left, top, width, Inches(4.6))
    first = True
    for item in items:
        if isinstance(item, tuple):
            head, subs = item
            para(tf, head, size=size, color=INK, bold=True, bullet=True, first=first, space_after=4)
            first = False
            for sub in subs:
                p = tf.add_paragraph()
                p.space_after = Pt(gap)
                p.level = 1
                r = p.add_run(); r.text = "–  " + sub
                r.font.size = Pt(sub_size); r.font.color.rgb = SOFT; r.font.name = FONT
        else:
            para(tf, item, size=size, color=INK, bullet=True, first=first, space_after=gap)
            first = False
    return tb


def picture_slide(kicker, title, subtitle, img, caption, n):
    s = add_slide()
    slide_title(s, kicker, title, subtitle)
    from PIL import Image
    im = Image.open(img)
    iw, ih = im.size
    max_w, max_h = Inches(11.9), Inches(4.55)
    ratio = min(max_w / iw, max_h / ih)
    w, h = int(iw * ratio), int(ih * ratio)
    left = (SW - w) // 2
    top = Inches(2.35)
    fr = rect(s, left - Pt(3), top - Pt(3), w + Pt(6), h + Pt(6), WHITE, line=True)
    s.shapes.add_picture(img, left, top, width=w, height=h)
    tb, tf = textbox(s, Inches(0.7), top + h + Inches(0.18), Inches(11.9), Inches(0.4))
    para(tf, caption, size=11.5, color=MUTED, italic=True, first=True)
    footer(s, n)
    return s


def stat_tile(s, l, t, w, h, value, label, color=RED):
    card = rrect(s, l, t, w, h, WHITE, line=True, radius=0.10)
    bar = rect(s, l, t, Pt(4), h, color)
    tb, tf = textbox(s, l + Inches(0.22), t + Inches(0.16), w - Inches(0.4), h - Inches(0.32),
                      valign=MSO_ANCHOR.MIDDLE)
    para(tf, value, size=26, color=INK, bold=True, first=True, space_after=2)
    para(tf, label, size=11.5, color=SOFT, space_after=0)


def phase_box(s, l, t, w, h, num, title, desc_lines, active=False):
    card = rrect(s, l, t, w, h, WHITE if not active else INK, line=True, radius=0.08)
    badge = s.shapes.add_shape(MSO_SHAPE.OVAL, l + Inches(0.18), t + Inches(0.16), Inches(0.4), Inches(0.4))
    badge.fill.solid(); badge.fill.fore_color.rgb = RED
    badge.line.fill.background(); badge.shadow.inherit = False
    btf = badge.text_frame; btf.margin_left = 0; btf.margin_right = 0
    bp = btf.paragraphs[0]; bp.alignment = PP_ALIGN.CENTER
    br = bp.add_run(); br.text = str(num); br.font.size = Pt(14); br.font.bold = True
    br.font.color.rgb = WHITE; br.font.name = FONT
    tb, tf = textbox(s, l + Inches(0.18), t + Inches(0.68), w - Inches(0.36), h - Inches(0.85))
    para(tf, title, size=13.5, color=(WHITE if active else INK), bold=True, first=True, space_after=5)
    for d in desc_lines:
        p = tf.add_paragraph(); p.space_after = Pt(3)
        r = p.add_run(); r.text = d
        r.font.size = Pt(10.5); r.font.color.rgb = (RGBColor(0xC8,0xC8,0xC8) if active else SOFT); r.font.name = FONT


def pipe_box(s, l, t, w, h, kicker, title, desc):
    card = rrect(s, l, t, w, h, WHITE, line=True, radius=0.10)
    tb, tf = textbox(s, l + Inches(0.15), t + Inches(0.14), w - Inches(0.3), h - Inches(0.28))
    para(tf, kicker.upper(), size=9.5, color=RED, bold=True, first=True, space_after=3)
    para(tf, title, size=13, color=INK, bold=True, space_after=4)
    para(tf, desc, size=10, color=SOFT, space_after=0)


def arrow(s, l, t, w, h):
    a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, l, t, w, h)
    a.fill.solid(); a.fill.fore_color.rgb = LINE
    a.line.fill.background(); a.shadow.inherit = False


def logo_placeholder(s, l, t, w, h, on_dark=False):
    """Marked, sized slot for the real DHL logo — paste it directly here in PowerPoint."""
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    try:
        box.adjustments[0] = 0.18
    except Exception:
        pass
    box.fill.background()
    box.line.color.rgb = (RGBColor(0x55, 0x55, 0x55) if on_dark else LINE)
    box.line.width = Pt(1.25)
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Pt(4); tf.margin_right = Pt(4)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = "Paste DHL logo here"
    r.font.size = Pt(9); r.font.italic = True
    r.font.color.rgb = (RGBColor(0x99, 0x99, 0x99) if on_dark else MUTED)
    r.font.name = FONT
    return box


# ============================================================ SLIDE 1 — TITLE
s = add_slide(bg=INK)
bar = rect(s, 0, 0, SW, Pt(10), YELLOW)
bar2 = rect(s, 0, Pt(10), SW, Pt(5), RED)
tb, tf = textbox(s, Inches(0.9), Inches(2.7), Inches(11), Inches(0.5))
para(tf, "DHL EXPRESS · AMKAD FINANCE", size=14, color=YELLOW, bold=True, first=True)
tb, tf = textbox(s, Inches(0.9), Inches(3.15), Inches(11.5), Inches(1.7))
para(tf, "AR Reporting Transformation Project", size=42, color=WHITE, bold=True, first=True, space_after=8)
tb, tf = textbox(s, Inches(0.9), Inches(4.55), Inches(11), Inches(0.6))
para(tf, "Internship Close-Out  ·  Jean Kadadihi  ·  DHL Express AMKAD Finance", size=16, color=RGBColor(0xC8,0xC8,0xC8), first=True)
rect(s, Inches(0.9), Inches(5.15), Inches(2.2), Pt(2.5), RED)
logo_placeholder(s, Inches(11.1), Inches(0.55), Inches(1.7), Inches(1.0), on_dark=True)

# ============================================================ SLIDE 2 — MISSION
s = add_slide()
slide_title(s, "Overview", "The Mission",
            "Move AMKAD AR reporting away from manual Excel processes toward a centralized, automated reporting ecosystem.")
bullets(s, [
    "Historical SAP data collection & standardization across countries",
    "Centralized AR database strategy",
    "Power BI executive dashboards",
    "Automated HTML executive reports (Office Scripts + Power Automate)",
    "Power Query ingestion automation",
    "Future SQL / Data Factory integration",
], size=16, gap=14)
tb, tf = textbox(s, Inches(0.7), Inches(6.15), Inches(11.9), Inches(0.9))
para(tf, "“Leading the technical side of an AMKAD AR reporting transformation — functioning as a Finance "
         "Analytics & Automation Developer rather than a typical intern.”", size=13, color=SOFT, italic=True, first=True)
footer(s, 2)

# ============================================================ SLIDE 3 — ROADMAP
s = add_slide()
slide_title(s, "Roadmap", "Five Phases of the Transformation")
phases = [
    (1, "SAP Data\nCollection", ["Historical AR extracts", "Open items data", "Country standardization"]),
    (2, "Centralized\nDatabase", ["Shared AR database", "Cross-country reporting", "Eliminate manual Excel"]),
    (3, "Power BI\nDevelopment", ["Executive KPIs", "Customer & aging views", "Replace monthly Excel"]),
    (4, "HTML Dashboard\nAutomation", ["Office Scripts + Power\nAutomate", "Fully automated pipeline", "Deepest technical work"]),
    (5, "SQL / Data\nFactory", ["Data Factory access\ngranted", "Centralized SQL layer", "Future integration"]),
]
bw = Inches(2.28); gap = Inches(0.16); startx = Inches(0.7); topy = Inches(2.5); bh = Inches(3.6)
for i, (num, title, desc) in enumerate(phases):
    x = startx + i * (bw + gap)
    phase_box(s, x, topy, bw, bh, num, title.replace("\n", " "), desc, active=(num == 4))
footer(s, 3)

# ============================================================ SLIDE 4 — FOUNDATIONS
s = add_slide()
slide_title(s, "Phase 1 & 2", "Foundations: SAP Data & Database Strategy")
tb, tf = textbox(s, Inches(0.7), Inches(2.35), Inches(5.8), Inches(4.3))
para(tf, "Phase 1 — Historical SAP Data Collection", size=15.5, color=RED, bold=True, first=True, space_after=8)
for t in ["Downloaded historical AR extracts & open items data",
          "Worked through SAP variants across countries",
          "Standardized country-level reporting formats",
          "Collaborated with Jose Pablo Cuevas, Sandra Garcia,\nJavier Leiva & country teams",
          "Built the historical dataset underpinning all\nlater automation"]:
    para(tf, t, size=13, color=INK, bullet=True, space_after=10)

tb, tf = textbox(s, Inches(6.85), Inches(2.35), Inches(5.75), Inches(4.3))
para(tf, "Phase 2 — Centralized AR Database", size=15.5, color=RED, bold=True, first=True, space_after=8)
for t in ["Strategy to upload AR files into a shared database",
          "Goal: centralized reporting across all countries",
          "Eliminate repetitive manual Excel processing",
          "Recurring governance meetings with Sandra\n(“Data Base & PBIs AMKAD”)",
          "Foundation for the SQL / Data Factory phase\nthat followed"]:
    para(tf, t, size=13, color=INK, bullet=True, space_after=10)
footer(s, 4)

# ============================================================ SLIDE 5 — POWER BI
s = add_slide()
slide_title(s, "Phase 3", "Power BI Development",
            "A documented implementation plan, aimed at replacing the existing monthly Excel report.")
bullets(s, [
    ("Executive Dashboard", ["Total AR, Overdue %, GT60 %, GT90 %"]),
    ("Customer View", ["Customer-level AR performance", "Top risk customers"]),
    ("Aging Analysis", ["Aging buckets & distribution analytics"]),
    ("Future Enhancements", ["Database connectivity", "Sales integration & DSO calculations", "Leadership-ready reporting"]),
], size=15.5, gap=8)
footer(s, 5)

# ============================================================ SLIDE 6 — PIPELINE
s = add_slide()
slide_title(s, "Phase 4", "HTML Dashboard Automation — The Pipeline",
            "The most technically complex piece of the project: a fully automated, self-refreshing reporting chain.")
stages = [
    ("Source", "Daily Excel File", "Uploaded to SharePoint each day"),
    ("Load", "Power Query", "Always finds the latest file"),
    ("Calculate", "Office Script", "MTD calculations, KPIs, tables"),
    ("Output", "JSON Result", "Structured data for the flow"),
    ("Assemble", "Power Automate", "Injects data into the template"),
    ("Deliver", "HTML Dashboard", "Emailed automatically"),
]
bw = Inches(1.78); gapw = Inches(0.12); arrw = Inches(0.28)
totalw = bw * 6 + gapw * 5 + arrw * 5
startx = (SW - totalw) // 2
topy = Inches(3.0); bh = Inches(1.7)
x = startx
for i, (k, t, d) in enumerate(stages):
    pipe_box(s, x, topy, bw, bh, k, t, d)
    x += bw + gapw
    if i < len(stages) - 1:
        arrow(s, x - gapw//2, topy + bh//2 - Inches(0.12), arrw, Inches(0.24))
        x += arrw
tb, tf = textbox(s, Inches(0.7), Inches(5.1), Inches(11.9), Inches(1.6))
para(tf, "Runs on a schedule with no manual intervention — already distributing “Weekly Automated Report” "
         "emails in production.", size=13.5, color=SOFT, italic=True, first=True)
footer(s, 6)

# ============================================================ SLIDE 7 — SNAPSHOT PROBLEM
s = add_slide()
slide_title(s, "Signature Technical Win", "Solving the “Monthly Snapshot” Problem",
            "The single most important concept in the reporting logic.")
tb, tf = textbox(s, Inches(0.7), Inches(2.35), Inches(6.7), Inches(4.4))
para(tf, "The problem:", size=14.5, color=INK, bold=True, first=True, space_after=6)
para(tf, "The AR source file contains multiple historical months and multiple daily snapshots per month, "
         "stretching back to 2024. Summing every row for a month meant a month's figures were counted "
         "20+ times over.", size=13, color=SOFT, space_after=14)
para(tf, "The fix:", size=14.5, color=INK, bold=True, space_after=6)
para(tf, "Keep only the latest valid snapshot for each month — a two-pass algorithm that detects the true "
         "point-in-time extract and discards every earlier daily copy.", size=13, color=SOFT, space_after=14)
para(tf, "Result: a Month-to-Date view that finance can actually trust.", size=13, color=INK, bold=True)

card = rrect(s, Inches(7.75), Inches(2.35), Inches(4.9), Inches(3.7), WHITE, line=True, radius=0.08)
tb, tf = textbox(s, Inches(8.0), Inches(2.55), Inches(4.4), Inches(3.3))
para(tf, "REAL IMPACT — BEFORE vs. AFTER", size=11, color=RED, bold=True, first=True, space_after=12)
para(tf, "June figures, before fix:", size=12.5, color=SOFT, space_after=2)
para(tf, "€924,030,792", size=22, color=CRIT, bold=True, space_after=12)
para(tf, "June figures, after fix:", size=12.5, color=SOFT, space_after=2)
para(tf, "€45,423,420", size=22, color=GOOD, bold=True, space_after=12)
para(tf, "~20× overstatement, caught and corrected", size=12.5, color=INK, bold=True)
footer(s, 7)

# ============================================================ SLIDES 8-11 — SCREENSHOTS
picture_slide("The Deliverable", "Executive Overview Dashboard",
              "Auto-generated narrative briefing, live KPIs, and a data-quality confidence banner.",
              SCRATCH + "f-top.png", "Live render of the deployed HTML report, current AMKAD data.", 8)

picture_slide("The Deliverable", "Portfolio Analytics",
              "AR concentration (Pareto/HHI) and a customer risk map — insights nobody had visibility into before.",
              SCRATCH + "f-portfolio.png", "Live render of the deployed HTML report, current AMKAD data.", 9)

picture_slide("The Deliverable", "Biggest Movers",
              "Automatically ranks the customers most improved and most deteriorated, month over month.",
              SCRATCH + "f-movers.png", "Live render of the deployed HTML report, current AMKAD data.", 10)

picture_slide("The Deliverable", "Trend Analysis",
              "Interactive, hover-enabled trend charts across 2+ years of point-in-time snapshots.",
              SCRATCH + "fix-trends-top.png", "Live render of the deployed HTML report, current AMKAD data.", 11)

# ============================================================ SLIDE 12 — TRUST & RELIABILITY
s = add_slide()
slide_title(s, "Governance", "Data Trust & Reliability",
            "A dashboard is only as good as the numbers behind it — built to fail loudly, never silently.")
bullets(s, [
    ("Confidence banner", ["Every report surfaces rows used vs. skipped, and any structural anomalies"]),
    ("Snapshot audit trail", ["Shows exactly which date was chosen for every month, fully auditable"]),
    ("Automated refresh reliability", ["A dedicated trigger flow guarantees the source file refreshes before the report runs"]),
    ("Freshness safeguards", ["Designed to fail visibly rather than silently ship stale numbers"]),
], size=15, gap=10)
footer(s, 12)

# ============================================================ SLIDE 13 — SQL / DATA FACTORY
s = add_slide()
slide_title(s, "Phase 5 — Looking Ahead", "SQL & Data Factory Integration")
bullets(s, [
    "Requested and received Data Factory access under Finance O2C reporting functions",
    "Reviewed SQL environments and centralization strategies",
    "Groundwork laid for moving from file-based automation to a true data platform",
    "Positions the next phase of this project for scale across AMKAD and beyond",
], size=16, gap=16)
footer(s, 13)

# ============================================================ SLIDE 14 — HOUSTON
s = add_slide()
slide_title(s, "Milestone", "Houston Working Session")
cols = [
    ("Power BI", ["Validated May numbers", "Reviewed charts & graphs", "Gross Sales review", "DSO logic review"]),
    ("SAP & Reporting", ["SAP download process\nimprovements", "HTML automation development", "Forecasting work", "SQL environment review"]),
    ("Process Exposure", ["Collections process exposure", "AMKAD overview", "Mexico SPR presentation\nobserved"]),
]
cw = Inches(3.85); gapx = Inches(0.16); startx = Inches(0.7); topy = Inches(2.4); ch = Inches(4.3)
x = startx
for title, items in cols:
    card = rrect(s, x, topy, cw, ch, WHITE, line=True, radius=0.06)
    tb, tf = textbox(s, x + Inches(0.2), topy + Inches(0.18), cw - Inches(0.4), ch - Inches(0.36))
    para(tf, title, size=14.5, color=RED, bold=True, first=True, space_after=10)
    for it in items:
        para(tf, it.replace("\n", " "), size=12, color=INK, bullet=True, space_after=9)
    x += cw + gapx
footer(s, 14)

# ============================================================ SLIDE 15 — PRODUCTION & OPS EXPOSURE
s = add_slide()
slide_title(s, "In Production", "Live Automation & Operational Exposure")
bullets(s, [
    ("Weekly Automated Reports", ["Already distributing production reports on a recurring schedule — development has moved to real usage"]),
    ("Statement of Account (SOA) shadowing", ["Shadowed Taralynn Haag on SOA workflows, invoice reconciliation, and Power BI usage in customer account reviews"]),
    ("Recurring governance", ["Regular Power BI plan review sessions with Sandra kept the project aligned with business needs"]),
], size=15.5, gap=12)
footer(s, 15)

# ============================================================ SLIDE 16 — CHALLENGES
s = add_slide()
slide_title(s, "Honest Retrospective", "Challenges Faced")
bullets(s, [
    "Diagnosing the multi-snapshot inflation bug — required building a diagnostic audit, not just guessing",
    "A month showing zero data (April) — traced to source-level gaps, not the automation itself",
    "Platform limits: Office Scripts cannot trigger a Power Query refresh — required a workaround flow design",
    "Power Automate silently stripping HTML on save — learned to always redeploy the full file, never a fragment",
    "Getting data-refresh timing right with no native “wait until ready” mechanism",
    "Standardizing AR business logic (Total AR, Overdue, GT60/90, DSO/TDSO) consistently across countries",
], size=15, gap=13)
footer(s, 16)

# ============================================================ SLIDE 17 — VALUE ADDED
s = add_slide()
slide_title(s, "Impact", "Value Added")
stats = [
    ("~20×", "Reporting distortion found\nand corrected", RED),
    ("100%", "Manual monthly Excel\nprocess automated", GOOD),
    ("2+ yrs", "Historical AR data\nstandardized", RED),
    ("Live", "Weekly automated reports\nin production", GOOD),
]
cw = Inches(2.78); gapx = Inches(0.2); startx = Inches(0.7); topy = Inches(2.5); ch = Inches(1.7)
x = startx
for val, lab, col in stats:
    stat_tile(s, x, topy, cw, ch, val, lab.replace("\n", " "), color=col)
    x += cw + gapx
bullets(s, [
    "Surfaced a real concentration risk invisible before: top customer ≈ 62% of AR, HHI ≈ 3,990",
    "Replaced a fully manual monthly reporting cycle with a self-refreshing automated pipeline",
    "Unlocked the path to SQL / Data Factory centralization for the next phase",
], top=Inches(4.6), size=14.5, gap=12)
footer(s, 17)

# ============================================================ SLIDE 18 — LEAVING BEHIND
s = add_slide()
slide_title(s, "Handover", "What I'm Leaving Behind")
bullets(s, [
    ("Maintenance guide, in three formats", ["Markdown (versioned in the repo), Word document, and a shareable web page"]),
    ("A documented Power BI implementation plan", ["KPIs, customer view, aging analysis, and a future-enhancements roadmap"]),
    ("A standardized historical SAP dataset", ["The foundation for all reporting going forward"]),
    ("A running, self-refreshing automation", ["Already in production, distributing weekly reports"]),
    ("Data Factory access under Finance O2C", ["A clear on-ramp to full SQL integration for my successor"]),
], size=14.5, gap=9)
footer(s, 18)

# ============================================================ SLIDE 19 — CLOSING
s = add_slide(bg=INK)
rect(s, 0, 0, SW, Pt(10), YELLOW)
rect(s, 0, Pt(10), SW, Pt(5), RED)
tb, tf = textbox(s, Inches(0.9), Inches(2.6), Inches(11.5), Inches(1.2))
para(tf, "Thank You", size=40, color=WHITE, bold=True, first=True)
tb, tf = textbox(s, Inches(0.9), Inches(3.75), Inches(10.8), Inches(2.5))
para(tf, "This internship took me from reading SAP extracts to designing a self-refreshing automated "
         "reporting platform — and taught me to build things that are still trustworthy after I'm gone.",
     size=15.5, color=RGBColor(0xD8,0xD8,0xD8), first=True, space_after=16)
para(tf, "Questions & discussion", size=14, color=YELLOW, bold=True)
logo_placeholder(s, Inches(11.1), Inches(0.55), Inches(1.7), Inches(1.0), on_dark=True)

prs.save(sys.argv[1] if len(sys.argv) > 1 else "AMKAD_AR_Transformation_Project.pptx")
print("Saved", prs.slides.__len__(), "slides")
