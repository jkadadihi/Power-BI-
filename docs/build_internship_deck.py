#!/usr/bin/env python3
"""
AMKAD Internship Close-Out deck.

Warm, humble, professional tone. No em dashes in slide text. Leads with a
transparent, conservative hours-saved model (~275-445 hrs/yr) built from the
intern's own figures, and tells the growth story from learning DHL's systems
in week 1 to shipping production automations by week 9.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------- palette ----------
RED = RGBColor(0xD4, 0x05, 0x11)
RED_DARK = RGBColor(0xA5, 0x04, 0x0D)
YELLOW = RGBColor(0xFF, 0xCC, 0x00)
INK = RGBColor(0x1A, 0x1D, 0x21)
INKCARD = RGBColor(0x23, 0x27, 0x2C)
SOFT = RGBColor(0x5B, 0x64, 0x70)
MUTED = RGBColor(0x8A, 0x83, 0x7B)
BG = RGBColor(0xF7, 0xF6, 0xF3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xE2, 0xE0, 0xDA)
GOOD = RGBColor(0x0F, 0x8A, 0x2E)
GOODBG = RGBColor(0xEC, 0xF6, 0xEE)
FONT = "Calibri"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height
PAGE = [0]


def add_slide(bg=BG):
    PAGE[0] += 1
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    r.fill.solid(); r.fill.fore_color.rgb = bg
    r.line.fill.background(); r.shadow.inherit = False
    s.shapes._spTree.remove(r._element)
    s.shapes._spTree.insert(2, r._element)
    return s


def rect(s, l, t, w, h, color, line=False):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    r.fill.solid(); r.fill.fore_color.rgb = color
    if line:
        r.line.color.rgb = LINE; r.line.width = Pt(0.75)
    else:
        r.line.fill.background()
    r.shadow.inherit = False
    return r


def rrect(s, l, t, w, h, color, line=False, radius=0.06):
    r = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    try:
        r.adjustments[0] = radius
    except Exception:
        pass
    r.fill.solid(); r.fill.fore_color.rgb = color
    if line:
        r.line.color.rgb = LINE; r.line.width = Pt(0.75)
    else:
        r.line.fill.background()
    r.shadow.inherit = False
    return r


def arrowshape(s, l, t, w, h, color):
    a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, l, t, w, h)
    a.fill.solid(); a.fill.fore_color.rgb = color
    a.line.fill.background(); a.shadow.inherit = False
    return a


def textbox(s, l, t, w, h, valign=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    return tb, tf


def para(tf, text, size=14, color=INK, bold=False, italic=False, align=PP_ALIGN.LEFT,
         space_after=6, first=False, bullet=False, font=FONT):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    if bullet:
        text = "•  " + text
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r.font.color.rgb = color; r.font.name = font
    return p


def eyebrow(s, text, top=Inches(0.5), left=Inches(0.7)):
    tb, tf = textbox(s, left, top, Inches(11), Inches(0.4))
    para(tf, text.upper(), size=12.5, color=RED, bold=True, first=True)
    rect(s, left, top + Inches(0.35), Inches(0.4), Pt(3.2), YELLOW)


def slide_title(s, kicker, title, subtitle=None):
    eyebrow(s, kicker)
    tb, tf = textbox(s, Inches(0.7), Inches(0.9), Inches(11.9), Inches(1.0))
    para(tf, title, size=29, color=INK, bold=True, first=True)
    if subtitle:
        tb2, tf2 = textbox(s, Inches(0.7), Inches(1.62), Inches(11.9), Inches(0.6))
        para(tf2, subtitle, size=14.5, color=SOFT, first=True)
    rect(s, Inches(0.7), Inches(2.12), Inches(11.9), Pt(1.4), LINE)


def footer(s):
    tb, tf = textbox(s, Inches(0.7), Inches(7.15), Inches(9), Inches(0.3))
    para(tf, "AMKAD  ·  AR Reporting & Collections Automation  ·  Internship Close-Out",
         size=9, color=MUTED, first=True)
    tb2, tf2 = textbox(s, Inches(12.3), Inches(7.15), Inches(0.7), Inches(0.3))
    para(tf2, str(PAGE[0]), size=9, color=MUTED, first=True, align=PP_ALIGN.RIGHT)


def logo_placeholder(s, l, t, w=Inches(1.9), h=Inches(0.75)):
    box = rrect(s, l, t, w, h, WHITE, line=True, radius=0.08)
    box.line.color.rgb = MUTED
    ln = box.line._get_or_add_ln()
    d = ln.makeelement(qn('a:prstDash'), {'val': 'dash'}); ln.append(d)
    tb, tf = textbox(s, l, t, w, h, valign=MSO_ANCHOR.MIDDLE)
    para(tf, "paste DHL logo", size=9, color=MUTED, italic=True, first=True, align=PP_ALIGN.CENTER)


def bullets(s, items, left=Inches(0.7), top=Inches(2.45), width=Inches(11.9),
            size=15, gap=9, color=INK):
    tb, tf = textbox(s, left, top, width, Inches(4.4))
    first = True
    for item in items:
        para(tf, item, size=size, color=color, bullet=True, first=first, space_after=gap)
        first = False
    return tb


def stat_tile(s, l, t, w, h, value, label, color=RED, vsize=30):
    rrect(s, l, t, w, h, WHITE, line=True, radius=0.10)
    rect(s, l, t, Pt(4), h, color)
    tb, tf = textbox(s, l + Inches(0.24), t + Inches(0.14), w - Inches(0.42), h - Inches(0.28),
                     valign=MSO_ANCHOR.MIDDLE)
    para(tf, value, size=vsize, color=INK, bold=True, first=True, space_after=2)
    para(tf, label, size=11.5, color=SOFT, space_after=0)


def card(s, l, t, w, h, title, lines, accent=RED, title_size=14, body_size=11.5):
    rrect(s, l, t, w, h, WHITE, line=True, radius=0.07)
    rect(s, l, t, w, Pt(4), accent)
    tb, tf = textbox(s, l + Inches(0.24), t + Inches(0.22), w - Inches(0.44), h - Inches(0.4))
    para(tf, title, size=title_size, color=INK, bold=True, first=True, space_after=7)
    for ln in lines:
        para(tf, ln, size=body_size, color=SOFT, bullet=True, space_after=5)


# =====================================================================
# TITLE
# =====================================================================
s = add_slide(INK)
rect(s, 0, 0, SW, Inches(0.18), RED)
rect(s, 0, Inches(0.18), SW, Pt(3), YELLOW)
logo_placeholder(s, Inches(10.9), Inches(0.55))
tb, tf = textbox(s, Inches(0.8), Inches(2.1), Inches(11), Inches(0.5))
para(tf, "DHL EXPRESS AMERICAS  ·  KEY ACCOUNT DESK (AMKAD)", size=13,
     color=YELLOW, bold=True, first=True)
tb, tf = textbox(s, Inches(0.8), Inches(2.7), Inches(11.7), Inches(2.0))
para(tf, "Automating AR Reporting & Collections", size=44, color=WHITE, bold=True, first=True,
     space_after=6)
para(tf, "Moving a few manual processes toward a live dashboard and pipelines that run on their own",
     size=17, color=RGBColor(0xC9, 0xCE, 0xD4))
tb, tf = textbox(s, Inches(0.8), Inches(5.7), Inches(11), Inches(1.0))
para(tf, "Jean Kadadihi", size=17, color=WHITE, bold=True, first=True, space_after=2)
para(tf, "Finance Intern, Controlling  ·  Internship Close-Out Presentation", size=13,
     color=RGBColor(0xB4, 0xBA, 0xC1))


# =====================================================================
# IMPACT HEADLINE
# =====================================================================
s = add_slide()
eyebrow(s, "The Impact")
tb, tf = textbox(s, Inches(0.7), Inches(0.9), Inches(11.9), Inches(1.0))
para(tf, "What the work added up to", size=29, color=INK, bold=True, first=True)
rect(s, Inches(0.7), Inches(1.75), Inches(11.9), Pt(1.4), LINE)
stat_tile(s, Inches(0.7), Inches(2.15), Inches(3.75), Inches(1.7),
          "~275 hrs", "of team time given back each year, and likely more. A conservative estimate.", vsize=32)
stat_tile(s, Inches(4.79), Inches(2.15), Inches(3.75), Inches(1.7),
          "6 / week", "statements now generated and emailed automatically (Arrow ×5, Baker ×1)", vsize=32)
stat_tile(s, Inches(8.88), Inches(2.15), Inches(3.72), Inches(1.7),
          "1 new", "shared AR reporting view the team did not have in one place before", vsize=32)
tb, tf = textbox(s, Inches(0.7), Inches(4.2), Inches(11.9), Inches(1.0))
para(tf, "That is roughly 7 to 11 full work-weeks a year that collectors can spend on collections instead of manual reporting.",
     size=17, color=RED_DARK, bold=True, first=True)
bullets(s, [
    "Helped bring AR visibility together in one place: aging, UAC, trends, and what-if scenarios.",
    "Took the daily data file off someone's plate so it no longer has to be rebuilt by hand.",
    "Reduced the reliance on one person, and the small errors that manual copy and paste can introduce.",
], top=Inches(4.85), size=14, gap=8)
footer(s)


# =====================================================================
# AGENDA
# =====================================================================
s = add_slide()
slide_title(s, "Agenda", "What I'll walk through")
items = [
    ("My internship journey", "From learning the systems to shipping automations"),
    ("Where things stood", "The manual starting point"),
    ("What I worked on", "Three connected automations"),
    ("The impact", "Time saved, with the math shown"),
    ("The hard parts and what I learned", "The challenges along the way, and the growth"),
    ("What's next", "How this can grow from here"),
]
x = Inches(0.7)
for i, (h, d) in enumerate(items):
    t = Inches(2.42) + i * Inches(0.75)
    badge = s.shapes.add_shape(MSO_SHAPE.OVAL, x, t, Inches(0.5), Inches(0.5))
    badge.fill.solid(); badge.fill.fore_color.rgb = RED; badge.line.fill.background()
    badge.shadow.inherit = False
    bp = badge.text_frame.paragraphs[0]; bp.alignment = PP_ALIGN.CENTER
    br = bp.add_run(); br.text = str(i + 1); br.font.size = Pt(16); br.font.bold = True
    br.font.color.rgb = WHITE; br.font.name = FONT
    tb, tf = textbox(s, x + Inches(0.75), t, Inches(10.5), Inches(0.55), valign=MSO_ANCHOR.MIDDLE)
    para(tf, h, size=16, color=INK, bold=True, first=True, space_after=1)
    para(tf, d, size=12, color=SOFT)
footer(s)


# =====================================================================
# MY INTERNSHIP JOURNEY  (new)
# =====================================================================
s = add_slide()
slide_title(s, "My Internship Journey",
            "From learning DHL's systems in week one to shipping automations that run on their own")
stages = [
    ("Weeks 1–2", "Learned AR, collections, START, and AMKAD processes; reviewed the existing Power BI reporting; picked up SAP extracts and the reporting workflow."),
    ("Weeks 3–4", "Built the first AR dashboard prototype and began data modeling and reporting enhancements."),
    ("Weeks 5–6", "Developed the automated daily data pipeline, connecting Power Query, Office Scripts, Power Automate, and SharePoint."),
    ("Weeks 7–8", "Reverse-engineered the START SOA process and designed a reusable SOA automation framework."),
    ("Week 9", "Deployed 5 Arrow SOAs and 1 Baker Hughes SOA, and completed documentation and handoff."),
]
top = Inches(2.45)
rh = Inches(0.9)
rail_x = Inches(2.15)
rect(s, rail_x, top + Inches(0.1), Pt(2.2), rh * (len(stages) - 1) + Inches(0.3), LINE)
for i, (wk, txt) in enumerate(stages):
    ry = top + i * rh
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, rail_x - Inches(0.09), ry + Inches(0.14),
                             Inches(0.26), Inches(0.26))
    dot.fill.solid(); dot.fill.fore_color.rgb = RED; dot.line.color.rgb = WHITE
    dot.line.width = Pt(2); dot.shadow.inherit = False
    tb, tf = textbox(s, Inches(0.7), ry, Inches(1.3), Inches(0.6))
    para(tf, wk, size=14, color=RED_DARK, bold=True, first=True)
    tb, tf = textbox(s, Inches(2.55), ry, Inches(10.0), Inches(0.8), valign=MSO_ANCHOR.TOP)
    para(tf, txt, size=13.5, color=INK, first=True)
footer(s)


# =====================================================================
# WHERE THINGS STOOD
# =====================================================================
s = add_slide()
slide_title(s, "Where Things Stood", "Three manual processes, three sources of risk")
card(s, Inches(0.7), Inches(2.45), Inches(3.85), Inches(3.9), "AR Reporting",
     ["No consolidated daily view of receivables",
      "Aging, overdue, UAC lived across spreadsheets",
      "No trend, drill-down, or what-if capability",
      "Leadership lacked a single place to look"], accent=RED)
card(s, Inches(4.74), Inches(2.45), Inches(3.85), Inches(3.9), "Daily Data File",
     ["A collector rebuilt it every business day",
      "Copy, paste, calculate, and save a new file",
      "Around 30 to 35 min a day, and mostly one person",
      "Manual copy and paste can introduce small errors"], accent=RED)
card(s, Inches(8.78), Inches(2.45), Inches(3.82), Inches(3.9), "Statements (SOA)",
     ["Collectors generated SOAs in START by hand",
      "Then formatted, emailed, attached, and CC'd",
      "Repeated weekly, per customer and country",
      "Missed sends and delays pulled focus off collections"], accent=RED)
footer(s)


# =====================================================================
# WHAT I WORKED ON
# =====================================================================
s = add_slide()
slide_title(s, "What I Worked On", "Three connected automations across the AR workflow")
card(s, Inches(0.7), Inches(2.45), Inches(3.85), Inches(3.9), "1 · AR Dashboard",
     ["Live HTML report of the whole book",
      "Aging, UAC, gross sales, payments, trends",
      "Drill-downs and what-if payment simulators",
      "Visibility we did not have in one place"], accent=RED)
card(s, Inches(4.74), Inches(2.45), Inches(3.85), Inches(3.9), "2 · Daily Pipeline",
     ["Feeds the dashboard automatically",
      "Source file lands, data is extracted and added",
      "New dated file built with full history",
      "Runs on its own, every business day"], accent=RED_DARK)
card(s, Inches(8.78), Inches(2.45), Inches(3.82), Inches(3.9), "3 · SOA Automation",
     ["6 weekly statements generated automatically",
      "Emailed with attachments and stakeholders CC'd",
      "A configurable framework, not one-offs",
      "Collectors step in only for exceptions"], accent=INK)
footer(s)


# =====================================================================
# HOW IT FITS TOGETHER  (architecture, new)
# =====================================================================
s = add_slide()
slide_title(s, "How It Fits Together",
            "The path from the source systems to the dashboard and the automated statements")
arch = [
    ("SAP / START", "source data"),
    ("SharePoint", "files land here"),
    ("Power Query", "shapes the data"),
    ("Office Script", "reads & writes"),
    ("Power Automate", "runs it daily"),
    ("Dashboard & SOA emails", "what people see"),
]
box_w = Inches(1.72)
gap = Inches(0.28)
y = Inches(3.15)
h = Inches(1.35)
x0 = Inches(0.72)
for i, (name, cap) in enumerate(arch):
    bx = x0 + i * (box_w + gap)
    last = (i == len(arch) - 1)
    b = rrect(s, bx, y, box_w, h, RED if last else WHITE, line=True, radius=0.10)
    tb, tf = textbox(s, bx + Inches(0.08), y + Inches(0.12), box_w - Inches(0.16), h - Inches(0.24),
                     valign=MSO_ANCHOR.MIDDLE)
    para(tf, name, size=12.5, color=WHITE if last else INK, bold=True, first=True,
         align=PP_ALIGN.CENTER, space_after=3)
    para(tf, cap, size=10, color=RGBColor(0xF3, 0xD6, 0xD6) if last else SOFT,
         align=PP_ALIGN.CENTER)
    if not last:
        arrowshape(s, bx + box_w + Inches(0.02), y + h / 2 - Inches(0.11),
                   gap - Inches(0.04), Inches(0.22), RGBColor(0xC7, 0xC2, 0xBA))
box = rrect(s, Inches(0.72), Inches(5.2), Inches(11.9), Inches(0.9), GOODBG, line=True, radius=0.10)
tb, tf = textbox(s, Inches(1.0), Inches(5.3), Inches(11.3), Inches(0.7), valign=MSO_ANCHOR.MIDDLE)
para(tf, "In plain terms: the data flows in on its own, gets shaped and checked automatically, "
         "and comes out as a live dashboard and ready-to-send statements. No one has to run it by hand.",
     size=13, color=RGBColor(0x1F, 0x5C, 0x2E), bold=True, first=True)
footer(s)


# =====================================================================
# PILLAR 1: DASHBOARD
# =====================================================================
s = add_slide()
slide_title(s, "Pillar 1 · The AR Dashboard",
            "Aging, UAC, payments, gross sales, trends, and forecasts brought into a single view")
bullets(s, [
    "Portfolio KPIs: Total AR, Overdue %, GT60 and GT90 aging, UAC, gross sales, payments.",
    "Drill-downs both ways, from country to customer, to see what is driving each number.",
    "Trends over time and same-month year-over-year comparison.",
    "What-if payment simulators, at portfolio and customer level, with manual allocation across aging bands and UAC.",
    "Editable risk thresholds and an Action-Required list that follows the collections escalation process.",
], top=Inches(2.45), size=15, gap=11)
box = rrect(s, Inches(0.7), Inches(5.95), Inches(11.9), Inches(0.85), RGBColor(0xFC, 0xEF, 0xEF),
            line=True, radius=0.12)
tb, tf = textbox(s, Inches(1.0), Inches(6.05), Inches(11.3), Inches(0.65), valign=MSO_ANCHOR.MIDDLE)
para(tf, "It replaced the report tab that used to be updated by hand each day, and brought AR "
         "visibility together in one place that the team did not have before.",
         size=13.5, color=RED_DARK, bold=True, first=True)
footer(s)


# =====================================================================
# PILLAR 2: PIPELINE
# =====================================================================
s = add_slide()
slide_title(s, "Pillar 2 · The Daily Data Pipeline",
            "The manual daily file, now running on its own")
tb, tf = textbox(s, Inches(0.7), Inches(2.4), Inches(5.7), Inches(0.4))
para(tf, "Before:  around 30 to 35 min a day, by hand", size=13, color=RED, bold=True, first=True)
bullets(s, [
    "Open the emailed report and copy the data",
    "Paste into a template and calculate",
    "Copy specific columns into the workbook",
    "Update the report tab, then save a new file",
], top=Inches(2.85), left=Inches(0.7), width=Inches(5.6), size=13, gap=7, color=SOFT)
tb, tf = textbox(s, Inches(6.85), Inches(2.4), Inches(5.7), Inches(0.4))
para(tf, "After:  it runs on its own", size=13, color=GOOD, bold=True, first=True)
bullets(s, [
    "The file lands in the shared folder and the flow starts",
    "An Office Script reads and maps today's rows",
    "A new dated file is built, carrying full history",
    "Today's rows are added and the dashboard updates",
], top=Inches(2.85), left=Inches(6.85), width=Inches(5.6), size=13, gap=7, color=SOFT)
box = rrect(s, Inches(0.7), Inches(5.55), Inches(11.9), Inches(1.25), INKCARD, radius=0.08)
tb, tf = textbox(s, Inches(1.0), Inches(5.72), Inches(11.3), Inches(0.95), valign=MSO_ANCHOR.MIDDLE)
para(tf, "Built with  Power Query (M),  Office Scripts (TypeScript),  Power Automate,  and SharePoint",
     size=13.5, color=YELLOW, bold=True, first=True, space_after=5)
para(tf, "It also protects against duplicates and handles missing or incomplete data gracefully, such as month-start EUR gaps.",
     size=12.5, color=RGBColor(0xCF, 0xD4, 0xDA))
footer(s)


# =====================================================================
# PILLAR 3: SOA
# =====================================================================
s = add_slide()
slide_title(s, "Pillar 3 · SOA Automation",
            "Recurring statements generated and delivered without a collector touching them")
bullets(s, [
    "Automated 6 recurring weekly Statements of Account: Arrow Electronics across 5 countries, and Baker Hughes.",
    "Generated from START templates with the right account filters, then emailed automatically with attachments and internal stakeholders CC'd.",
    "Built as a configurable framework, so new customers can be added by updating configuration (accounts, recipients, schedule, template) rather than rebuilding it.",
    "Designed so collectors only step in for exceptions, and can spend more of their time on collections.",
], top=Inches(2.45), size=14.5, gap=11)
box = rrect(s, Inches(0.7), Inches(5.95), Inches(11.9), Inches(0.85), RGBColor(0xFC, 0xEF, 0xEF),
            line=True, radius=0.12)
tb, tf = textbox(s, Inches(1.0), Inches(6.05), Inches(11.3), Inches(0.65), valign=MSO_ANCHOR.MIDDLE)
para(tf, "To be clear on scope: two customers are live today, six SOAs a week, on a framework built so more "
         "can be added over time. The reusable engine is the part that is done.",
         size=13, color=RED_DARK, bold=True, first=True)
footer(s)


# =====================================================================
# THE IMPACT (table + current state)
# =====================================================================
s = add_slide()
slide_title(s, "The Impact", "Time saved, with the math shown, kept deliberately conservative")

# current-state strip
cs = rrect(s, Inches(0.7), Inches(2.28), Inches(11.9), Inches(0.5), GOODBG, line=True, radius=0.16)
tb, tf = textbox(s, Inches(1.0), Inches(2.28), Inches(11.3), Inches(0.5), valign=MSO_ANCHOR.MIDDLE)
para(tf, "Live today:   ✓ Dashboard in production    ✓ Daily pipeline running    "
         "✓ 6 weekly SOAs automated    ✓ Framework ready for more customers",
     size=12.5, color=RGBColor(0x1F, 0x5C, 0x2E), bold=True, first=True)

rows = [
    ("", "Per week", "Per month", "Per year", False),
    ("Daily data file  (15 min × 5 days)", "1.25 hrs", "~5.4 hrs", "~63 hrs", False),
    ("Daily report tab  (15–20 min × 5 days)", "1.25–1.7 hrs", "~5–7 hrs", "~63–83 hrs", False),
    ("SOA automation  (6 SOAs × 30–60 min)", "3–6 hrs", "13–26 hrs", "150–300 hrs", False),
    ("Combined", "~5.5–9 hrs", "~24–38 hrs", "~275–445 hrs", True),
]
tw = Inches(11.9)
c0, c1, c2, c3 = Inches(5.0), Inches(2.3), Inches(2.3), Inches(2.3)
x0 = Inches(0.7)
y = Inches(2.95)
rh = Inches(0.55)
for i, (label, a, b, c, total) in enumerate(rows):
    ry = y + i * rh
    if i == 0:
        rect(s, x0, ry, tw, rh, INK)
        cells = [(c0, label, PP_ALIGN.LEFT), (c1, a, PP_ALIGN.CENTER),
                 (c2, b, PP_ALIGN.CENTER), (c3, c, PP_ALIGN.CENTER)]
        cx = x0
        for cw, txt, al in cells:
            tb, tf = textbox(s, cx + Inches(0.15), ry, cw - Inches(0.3), rh, valign=MSO_ANCHOR.MIDDLE)
            para(tf, txt, size=13, color=YELLOW, bold=True, first=True, align=al)
            cx += cw
    else:
        fill = RGBColor(0xFC, 0xEF, 0xEF) if total else (WHITE if i % 2 else RGBColor(0xF1, 0xEF, 0xEA))
        rect(s, x0, ry, tw, rh, fill, line=True)
        if total:
            rect(s, x0, ry, Pt(4), rh, RED)
        cells = [(c0, label, PP_ALIGN.LEFT), (c1, a, PP_ALIGN.CENTER),
                 (c2, b, PP_ALIGN.CENTER), (c3, c, PP_ALIGN.CENTER)]
        cx = x0
        for j, (cw, txt, al) in enumerate(cells):
            tb, tf = textbox(s, cx + Inches(0.15), ry, cw - Inches(0.3), rh, valign=MSO_ANCHOR.MIDDLE)
            para(tf, txt, size=13.5 if total else 13,
                 color=RED_DARK if total else INK, bold=(total or j == 0), first=True, align=al)
            cx += cw

tb, tf = textbox(s, Inches(0.7), Inches(5.95), Inches(11.9), Inches(1.2))
para(tf, "That comes to roughly 7 to 11 full work-weeks of time given back to the team each year.",
     size=15.5, color=RED_DARK, bold=True, first=True, space_after=7)
para(tf, "The assumptions are kept conservative on purpose: 5 business days a week, SOAs at 30 to 60 minutes each, "
         "6 a week. The figures leave out error-correction, rework, and analysis time, so the real number is likely higher.",
     size=11, color=MUTED, italic=True)
footer(s)


# =====================================================================
# THE HARD PARTS
# =====================================================================
s = add_slide()
slide_title(s, "The Hard Parts", "What it took along the way, and what I learned")
card(s, Inches(0.7), Inches(2.4), Inches(3.85), Inches(2.05), "Power BI & the data",
     ["A lot of trial and error on the model",
      "Fragmented, inconsistent source data",
      "Had to consolidate before trusting anything"], accent=RED, title_size=13.5)
card(s, Inches(4.74), Inches(2.4), Inches(3.85), Inches(2.05), "Hosting blocked",
     ["Azure Functions blocked by permissions",
      "Rethought the recurring-run approach",
      "Landed on GitHub Actions, clean and no cost"], accent=RED, title_size=13.5)
card(s, Inches(8.78), Inches(2.4), Inches(3.82), Inches(2.05), "The right data pattern",
     ["Power Query replaces, it cannot append",
      "Designed a carry-history-forward pipeline",
      "Solved it without breaking the live report"], accent=RED, title_size=13.5)
card(s, Inches(0.7), Inches(4.6), Inches(3.85), Inches(2.05), "A production incident",
     ["A live query edit went wrong",
      "Recovered cleanly via version history",
      "Adopted a test-first, never-in-prod habit"], accent=INK, title_size=13.5)
card(s, Inches(4.74), Inches(4.6), Inches(3.85), Inches(2.05), "Connector quirks",
     ["Many field-name and identifier mismatches",
      "Pagination, throttling, locked-file retries",
      "Worked each one through to a clean run"], accent=INK, title_size=13.5)
card(s, Inches(8.78), Inches(4.6), Inches(3.82), Inches(2.05), "The takeaway",
     ["Kept going through the obstacles",
      "Learned to de-risk before touching live data",
      "Left it documented and maintainable"], accent=INK, title_size=13.5)
footer(s)


# =====================================================================
# THE BIGGER CHALLENGE  (reflective challenges, new)
# =====================================================================
s = add_slide()
slide_title(s, "The Bigger Challenge Wasn't the Code",
            "The hardest parts were the ones around the building")
bullets(s, [
    "Learning the business first: AR, aging, UAC, DSO, SOAs, START, and the AMKAD workflows, before automating any of it.",
    "Understanding why collectors did each step, what the exceptions were, and what could break if I got it wrong.",
    "Making imperfect, fragmented data reliable: different files, SAP exports, missing values, and inconsistent naming.",
    "Building systems people actually depend on, where accuracy and uptime matter, not demo projects.",
    "Piecing together processes that were never fully documented, by meeting people, asking questions, and testing.",
], top=Inches(2.45), size=15, gap=12)
footer(s)


# =====================================================================
# WHAT I LEARNED  (reflection, new)
# =====================================================================
s = add_slide()
slide_title(s, "What I Learned", "The parts that will stay with me")
bullets(s, [
    "How accounts receivable connects to the wider health and performance of the business.",
    "How to take a business problem and turn it into a practical automation opportunity.",
    "How to work with people across finance, collections, and IT to get something built.",
    "Why it is worth designing scalable, reusable solutions instead of one-off fixes.",
    "How to stay calm and methodical when something breaks, and recover from it properly.",
], top=Inches(2.5), size=15.5, gap=13)
footer(s)


# =====================================================================
# WHAT I'D DO DIFFERENTLY  (self-reflection, new)
# =====================================================================
s = add_slide()
slide_title(s, "What I'd Do Differently", "Where I want to grow from here")
bullets(s, [
    "Spend more time on design up front. Map the process and sketch the architecture before building, so there is less rebuilding later when a new requirement appears.",
    "Document while I build, not at the end. As the automations grew more complex, earlier notes would have made them easier to hand off and maintain.",
    "Validate the workflow with users earlier and more often. Not just \"can I make this work,\" but \"is this the workflow people actually want.\"",
], top=Inches(2.5), size=15.5, gap=15)
box = rrect(s, Inches(0.7), Inches(5.75), Inches(11.9), Inches(0.9), RGBColor(0xFC, 0xEF, 0xEF),
            line=True, radius=0.10)
tb, tf = textbox(s, Inches(1.0), Inches(5.85), Inches(11.3), Inches(0.7), valign=MSO_ANCHOR.MIDDLE)
para(tf, "In short, I want to grow from building automations toward designing them: more thought up front, "
         "better documentation, and scalability from the start.", size=13.5, color=RED_DARK, bold=True, first=True)
footer(s)


# =====================================================================
# SKILLS
# =====================================================================
s = add_slide()
slide_title(s, "What It Took", "Skills I got to build and use")
groups = [
    ("Automation & Data", ["Power Query (M)", "Office Scripts (TypeScript)",
                            "Power Automate", "SharePoint / GitHub Actions"]),
    ("Analysis & Modeling", ["AR aging & cash-application logic", "Data consolidation & modeling",
                             "Power BI", "MTD / snapshot filtering"]),
    ("Product & Front-End", ["HTML / CSS / JavaScript", "SVG charting & dashboards",
                             "UX for non-technical users", "Clear documentation"]),
    ("Ways of Working", ["Reverse-engineering undocumented processes", "Working from ambiguity",
                         "Production troubleshooting", "Following things through to done"]),
]
for i, (title, items) in enumerate(groups):
    col = i % 2
    row = i // 2
    l = Inches(0.7) + col * Inches(6.1)
    t = Inches(2.5) + row * Inches(2.25)
    rrect(s, l, t, Inches(5.85), Inches(2.0), WHITE, line=True, radius=0.07)
    rect(s, l, t, Pt(4), Inches(2.0), RED)
    tb, tf = textbox(s, l + Inches(0.25), t + Inches(0.18), Inches(5.4), Inches(1.7))
    para(tf, title, size=14, color=RED_DARK, bold=True, first=True, space_after=7)
    for it in items:
        para(tf, it, size=12.5, color=INK, bullet=True, space_after=4)
footer(s)


# =====================================================================
# WHAT I'M MOST PROUD OF  (new)
# =====================================================================
s = add_slide()
slide_title(s, "What I'm Most Proud Of", "A few things that meant the most to me")
bullets(s, [
    "Building three automation solutions that are genuinely in use, not just prototypes.",
    "Taking real, repetitive reporting effort off the team's plate each day.",
    "Creating reusable frameworks that can grow, rather than one-time fixes.",
    "Leaving behind solutions that keep running after the internship ends.",
    "Giving the team clearer visibility into AR performance and collections activity.",
], top=Inches(2.5), size=15.5, gap=13)
footer(s)


# =====================================================================
# WHAT'S NEXT
# =====================================================================
s = add_slide()
slide_title(s, "What's Next", "How this can grow from here")
bullets(s, [
    "Extend the SOA framework to more customers and countries, since the reusable engine is already in place.",
    "Move the data layer toward a centralized database (SQL and Data Factory), which would be a stronger backbone for everything above.",
    "Build leadership-facing Power BI dashboards on top of the same consolidated data.",
    "Small hardening steps, like standardizing the monthly folder names and adding alerts for missing or late source files.",
], top=Inches(2.5), size=15.5, gap=13)
footer(s)


# =====================================================================
# LOOKING AHEAD (close)
# =====================================================================
s = add_slide(INK)
rect(s, 0, 0, SW, Inches(0.16), RED)
rect(s, 0, Inches(0.16), SW, Pt(3), YELLOW)
tb, tf = textbox(s, Inches(0.8), Inches(0.75), Inches(11), Inches(0.5))
para(tf, "LOOKING AHEAD", size=13, color=YELLOW, bold=True, first=True)
tb, tf = textbox(s, Inches(0.8), Inches(1.35), Inches(11.7), Inches(1.0))
para(tf, "I have genuinely loved being part of this team.", size=30, color=WHITE, bold=True, first=True)
pts = [
    "Thank you for the trust, the patience, and the chance to work on real problems that matter to the team.",
    "I learned an enormous amount here, from the AR and collections side of the business to the tools that support it.",
    "I am proud of what we built together, and I know there is more I could contribute with more time.",
    "This internship confirmed my interest in finance, analytics, and automation.",
]
tb, tf = textbox(s, Inches(0.8), Inches(2.7), Inches(11.7), Inches(3.2))
first = True
for p in pts:
    para(tf, p, size=16, color=RGBColor(0xE7, 0xEA, 0xED), bullet=True, first=first, space_after=15)
    first = False
box = rrect(s, Inches(0.8), Inches(6.05), Inches(11.7), Inches(0.85), RED, radius=0.12)
tb, tf = textbox(s, Inches(1.1), Inches(6.14), Inches(11.1), Inches(0.65), valign=MSO_ANCHOR.MIDDLE)
para(tf, "I am grateful for the opportunity to contribute to AMKAD, and would be excited to keep "
         "building solutions like these in the future.", size=15, color=WHITE, bold=True, first=True)


# =====================================================================
# THANK YOU
# =====================================================================
s = add_slide()
rect(s, 0, Inches(3.5), SW, Pt(3), RED)
tb, tf = textbox(s, Inches(0.8), Inches(2.5), Inches(11.7), Inches(1.0), valign=MSO_ANCHOR.BOTTOM)
para(tf, "Thank you", size=40, color=INK, bold=True, first=True)
tb, tf = textbox(s, Inches(0.8), Inches(3.75), Inches(11.7), Inches(1.2))
para(tf, "Jean Kadadihi", size=18, color=INK, bold=True, first=True, space_after=3)
para(tf, "Finance Intern, Controlling  ·  DHL Express Americas, Key Account Desk (AMKAD)",
     size=13.5, color=SOFT, space_after=2)
para(tf, "I would be happy to answer any questions.", size=13.5, color=RED_DARK, italic=True)
logo_placeholder(s, Inches(10.9), Inches(0.6))


out = "docs/AMKAD_Internship_Presentation.pptx"
prs.save(out)
print("Saved", out, "-", len(prs.slides._sldIdLst), "slides")
