#!/usr/bin/env python3
"""
Second update pass on the AMKAD HTML Maintenance Guide (.docx), IN PLACE.

Run this AFTER update_maintenance_docx.py has already produced a guide that
contains sections 5F (daily file build) and 5G (SPR tab). It edits that
document rather than regenerating it, because the file carries 7 embedded
screenshots that no Markdown source knows about — a rebuild would drop them
silently. Every new element is cloned from one already in the document, so
fonts, borders and shading match without hardcoding any styling.

Covers the work done after the first pass:
  - Section 5H: the self-healing RawData query that repairs days lost to a
    failed pipeline run
  - Section 5G expanded: country-first drilldown, dual month-end/current
    figures with source labels, GT60/GT90 %, GT90 sort order, country-level
    over-90 %
  - Section 5B expanded: why a fall is green for some metrics and red for
    others, why DSO/TDSO are excluded, and MoM vs YoY
  - Section 5A: the Simulator now allocates to UAC and not-yet-overdue AR
  - Action Required: the escalation stages are approximations and say so
  - New troubleshooting entries and maintenance-calendar rows

No repository paths appear anywhere. Whoever inherits this will not have the
Git repo — it was only ever our working copy on a personal account — so every
component is named by where they open it.

Usage: python3 docs/update_maintenance_docx_v2.py <in.docx> <out.docx>
"""
import copy
import sys

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph


# ---------- helpers that clone existing formatting ----------

def set_cell_text(cell, text, bold=None):
    """Replace a cell's text while keeping its first run's formatting."""
    paras = cell.paragraphs
    first = paras[0]
    for extra in paras[1:]:
        extra._element.getparent().remove(extra._element)
    if first.runs:
        keep = first.runs[0]
        for r in first.runs[1:]:
            r._element.getparent().remove(r._element)
        keep.text = text
        if bold is not None:
            keep.font.bold = bold
    else:
        run = first.add_run(text)
        if bold is not None:
            run.font.bold = bold


def clone_row(table, source_index, values):
    """Append a row copied from an existing one, so borders/shading match."""
    src = table.rows[source_index]._tr
    new = copy.deepcopy(src)
    table._tbl.append(new)
    row = table.rows[-1]
    for cell, val in zip(row.cells, values):
        set_cell_text(cell, val)
    return row


def clone_callout(doc, template_table, text, before):
    """Copy a 1x1 callout box and drop in new text."""
    new = copy.deepcopy(template_table._tbl)
    before.addprevious(new)
    tbl = Table(new, doc)
    set_cell_text(tbl.rows[0].cells[0], text)
    return tbl


def clone_para(doc, template_para, text, before):
    """Copy a paragraph's styling and insert it before `before`."""
    new = copy.deepcopy(template_para._p)
    before.addprevious(new)
    p = Paragraph(new, doc)
    if p.runs:
        keep = p.runs[0]
        for r in p.runs[1:]:
            r._element.getparent().remove(r._element)
        keep.text = text
    else:
        p.add_run(text)
    return p


def find_para(doc, predicate):
    for p in doc.paragraphs:
        if predicate(p.text.strip()):
            return p
    return None


def find_table(doc, predicate):
    """Locate a table by its header text, not by index.

    The first pass inserted callout tables, which shifts every index after
    them. Matching on content survives that.
    """
    for t in doc.tables:
        if predicate([c.text.strip() for c in t.rows[0].cells]):
            return t
    return None


def main(src_path, out_path):
    doc = Document(src_path)

    images_before = len(doc.inline_shapes)

    # Guard: this pass assumes the first one already ran.
    if find_para(doc, lambda t: t.startswith("F · The daily file build")) is None:
        sys.exit("This document has no Section 5F. Run update_maintenance_docx.py first.")
    if find_para(doc, lambda t: t.startswith("H · Repairing")) is not None:
        sys.exit("This document already has Section 5H — it looks like this pass "
                 "has already been applied. Re-run from the previous version.")

    # Style templates lifted from the document itself.
    tmpl_h3 = next(p for p in doc.paragraphs if p.style.name == "Heading 3")
    tmpl_body = find_para(doc, lambda t: t.startswith("Where it lives:"))
    tmpl_list = next(p for p in doc.paragraphs if p.style.name == "List Number")
    tmpl_callout = find_table(doc, lambda h: h and h[0].startswith("Naming."))

    # ---------- 1 · components table ----------
    t_files = find_table(doc, lambda h: h[:2] == ["COMPONENT", "WHAT IT IS"])
    clone_row(t_files, 1, (
        "Power Query query: RawData",
        "Reads the newest daily file and repairs any days missing from it (5H)",
        "Excel → Power Query",
    ))

    # ---------- 2 · maintenance calendar ----------
    t_cal = find_table(doc, lambda h: h[:2] == ["WHEN", "DO THIS"])
    for vals in [
        ("Weekly",
         "Glance at the row count after a RawData refresh. It should grow by roughly one "
         "day's worth of rows per working day. A flat count means the pipeline stopped "
         "producing new files and nobody was told."),
        ("Whenever a run fails",
         "Nothing, if the self-healing RawData query is in use — the missing day is picked "
         "up from the source file on the next refresh (5H). Confirm it actually came back "
         "before assuming it did."),
    ]:
        clone_row(t_cal, 1, vals)

    # ---------- 3 · additions to 5G, then the new 5H ----------
    anchor6 = find_para(doc, lambda t: t == "SECTION 06")._p

    def h3(text):
        clone_para(doc, tmpl_h3, text, anchor6)

    def body(text):
        clone_para(doc, tmpl_body, text, anchor6)

    def item(text):
        clone_para(doc, tmpl_list, text, anchor6)

    def callout(text):
        clone_callout(doc, tmpl_callout, text, anchor6)

    # -- 5G continued: how the tab is laid out now
    body("How the tab is organised. Reviews are run country by country, so the tab opens "
         "on a list of countries rather than a list of accounts. Choosing one expands its "
         "customers, each of which opens individually. Customers are ordered by their "
         "over-90 amount, largest first, so the accounts that will take up the meeting are "
         "at the top without scrolling.")
    body("Each country strip also carries its own over-90 percentage. This is calculated "
         "from the country's totals — the sum of over-90 divided by the sum of AR — and is "
         "NOT the average of the customer percentages underneath it. Averaging those would "
         "weight a small account the same as a large one and produce a figure that matches "
         "nothing.")
    callout("Which numbers you are looking at.  Every set of figures on this tab is labelled "
            "with the table it came from: month-end close, or current/daily. The two are "
            "shown one above the other for the same account so they can be compared, and "
            "the country totals declare a single basis and stick to it. If a month-end "
            "figure is shown for a different period than the commentary — normal mid-month, "
            "when the current month has not closed yet — the tab names both periods rather "
            "than quietly substituting one for the other.")

    # -- 5H: the self-healing query
    h3("H · Repairing days lost to a failed run")
    body("Where it lives: the input workbook → Data → Queries → RawData.")
    body("Each daily file is built by copying the previous day's file and adding one day to "
         "it. That design carries the history forward cheaply, but it has one consequence "
         "worth understanding: if a run ever fails, that day is absent from the file built "
         "the next day, and from every file built after that. No later refresh brings it "
         "back on its own, because the newest file is the only one the report reads.")
    body("The raw source files, however, are all still in the Cust Sol folder. The RawData "
         "query reads the newest daily file as before, then checks whether any source file "
         "has a date the history does not contain, and fills those days in from the source.")
    item("It only ever adds rows. It cannot remove or alter a day that is already there.")
    item("It only fills gaps INSIDE the range it already holds. Source files older than the "
         "first date in the history are ignored, so it repairs holes rather than silently "
         "extending the history backwards.")
    item("On a normal day nothing is missing, so the repair does not run and the refresh "
         "costs what it always did.")
    item("Backfilled days have no TDSO, DSO or percentage values. Those are formulas that "
         "live in the daily file, and a day recovered from the source never passed through "
         "one. The amounts are complete; the calculated columns are blank for those days.")
    callout("If the refresh reports a download failure.  The repair opens one extra file per "
            "missing day, so it makes more calls to SharePoint than the old single-file "
            "version did and is more exposed to a transient failure. Each file is handled "
            "independently: one that cannot be opened costs that day only, and if the whole "
            "repair fails the query still returns the history unchanged. Refresh again "
            "before investigating — most of these clear on their own.")

    # ---------- 4 · reading the report correctly (into 5B) ----------
    anchor_c = find_para(doc, lambda t: t.startswith("C · Which daily file"))._p

    def body_b(text):
        clone_para(doc, tmpl_body, text, anchor_c)

    def callout_b(text):
        clone_callout(doc, tmpl_callout, text, anchor_c)

    callout_b("Arrows and colours.  A movement is coloured by whether it is GOOD, not by "
              "which way it points. A fall in AR, overdue, over-60, over-90 or unapplied "
              "cash is green; a rise in payments or sales is green. If you add a metric, "
              "say which direction is favourable for it — the report does not infer this, "
              "and getting it wrong produces a red arrow on good news, which is worse than "
              "no arrow at all.")
    body_b("Aging buckets are nested, not separate slices: over-90 is part of over-60, which "
           "is part of overdue, which is part of total AR. They must never be stacked or put "
           "in a pie chart, because doing so counts the same money more than once. Side-by-"
           "side bars are the correct way to show them together.")
    body_b("DSO and TDSO are read from the source but deliberately excluded from every "
           "calculated figure. They are ratios, and rolling them up by averaging customer "
           "values would weight a small account the same as a large one. A meaningful "
           "portfolio or country DSO has to be recalculated from the underlying totals, "
           "which is a change to make deliberately rather than by leaving the columns in.")
    body_b("On the Customers tab, the two aging tables compare the three most recent "
           "consecutive months — that is month-over-month. The same-month year-over-year "
           "comparison is the portfolio-level table on the Trends tab, which compares the "
           "same calendar month across three years. They answer different questions and are "
           "easy to mix up when reading quickly.")

    # ---------- 5 · the Simulator ----------
    tbl_sim = find_table(doc, lambda h: h and h[0].startswith("The Simulator tab is calibra"))
    if tbl_sim is not None:
        nxt = tbl_sim._tbl.getnext()
        if nxt is not None:
            clone_callout(
                doc, tmpl_callout,
                "Allocating cash in the Simulator.  Payments can be applied to unapplied "
                "cash and to AR that is not yet overdue, as well as to the over-60 and "
                "over-90 buckets. This is what makes the over-90 percentage behave in a way "
                "that surprises people: clearing cash off a younger bucket lowers total AR "
                "without lowering the over-90 amount, so over-90 as a PERCENTAGE goes up "
                "even though the situation improved. The amount is the figure to watch when "
                "judging whether an allocation helped.",
                nxt)

    # ---------- 6 · Action Required ----------
    anchor_d = find_para(doc, lambda t: t.startswith("D · The automated email"))
    if anchor_d is not None:
        clone_callout(
            doc, tmpl_callout,
            "The Action Required tab.  Suggested next steps are mapped from the aging "
            "buckets onto the collections escalation stages. This is an APPROXIMATION: the "
            "report knows how old a balance is, not why, and reason matters as much as age "
            "when deciding what to do. Every suggestion is therefore worded as a potential "
            "step and needs a person's review before anything is actioned. No step that "
            "would hold or stop a customer's service is ever suggested automatically.",
            anchor_d._p)

    # ---------- 7 · troubleshooting ----------
    anchor7 = find_para(doc, lambda t: t == "SECTION 07")._p
    for text in [
        "ROW COUNT JUMPED   RawData suddenly returned noticeably more rows than last time.  "
        "Expected the first time the self-healing query runs: it recovers every day that "
        "was missing, so the increase should be a whole number of days' worth of rows. "
        "Check the DATES that appeared rather than the count — they should be days that "
        "were genuinely absent, and all within the range the history already covered. If "
        "new dates appear OLDER than the previous earliest date, something is wrong; the "
        "query is meant to fill holes, not extend the history backwards.",

        "PERCENTAGES BLANK ON SOME DAYS   A few days have amounts but no TDSO, DSO or "
        "percentage values.  These are days recovered from the source file by the repair "
        "described in 5H. Those columns are formulas in the daily file, and a recovered day "
        "never passed through it. The amounts are correct and complete; nothing needs "
        "fixing.",

        "OVER-90 % ROSE AFTER A PAYMENT   An allocation was posted and the percentage went "
        "up.  Almost always correct behaviour, not a bug. The percentage is over-90 divided "
        "by total AR, so applying cash to a younger bucket shrinks the denominator while "
        "the over-90 amount stays put. Compare the over-90 AMOUNT before and after to see "
        "whether the position actually improved.",
    ]:
        clone_callout(doc, tmpl_callout, text, anchor7)

    # ---------- 8 · sweep up bare filenames the first pass missed ----------
    # The first pass rewrote path-prefixed names ("power_query/spr_month_end_ar.m")
    # but not bare ones written mid-sentence. Nobody inheriting this can open a
    # file by that name, so name the thing they actually click instead.
    bare_renames = {
        # The first pass swapped the path for a friendly name but left the
        # "File:" lead-in, which now reads oddly since none of these are files
        # the reader can open.
        "File: HTML template (report flow)":
            "Where it lives: the HTML / Compose action inside the report flow",
        "File: Office Script: AMKAD AR report":
            "Where it lives: Excel Online → Automate → the AMKAD AR report script",
        "File: Email body (Send an email V2)":
            "Where it lives: Power Automate → Send an email (V2) → Body",
        "File: Office Script: refresh trigger":
            "Where it lives: Excel Online → Automate → the refresh-trigger script",
        "spr_month_end_ar.m": "the MonthEndAR query",
        "custsol_rawdata.m": "the RawData query",
        "amkad_ar_report.ts": "the AMKAD AR report script",
        "custsol_read_daily.ts": "the 'custsol read daily' script",
        "custsol_append_rawdata_eur.ts": "the 'custsol append rawdata eur' script",
        "refresh_trigger.ts": "the refresh-trigger script",
        "html_template.html": "the HTML template",
        "email_body.html": "the email body",
    }

    def rewrite_para(p):
        original = p.text
        updated = original
        for old, new in bare_renames.items():
            updated = updated.replace(old, new)
        if updated == original or not p.runs:
            return
        keep = p.runs[0]
        for r in p.runs[1:]:
            r._element.getparent().remove(r._element)
        keep.text = updated

    for p in doc.paragraphs:
        rewrite_para(p)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    rewrite_para(p)

    # ---------- 9 · verify nothing was lost ----------
    images_after = len(doc.inline_shapes)
    if images_after != images_before:
        sys.exit("Image count changed (%d -> %d) — aborting rather than saving a guide "
                 "with missing screenshots." % (images_before, images_after))

    doc.save(out_path)
    print("Saved %s  (screenshots preserved: %d)" % (out_path, images_after))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
