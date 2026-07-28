#!/usr/bin/env python3
"""
Generate the SPR (Statement Performance Review) commentary sheets, one worksheet
per country, pre-filled with that country's customers so collectors only have to
type the commentary and the forecast numbers.

Each country sheet holds two Excel Tables:

  SPRIssues_<CC>    Month | Country | Customer | Issue | Action
                    One row per issue, so nobody has to cram four paragraphs
                    into a single cell. Pre-seeded with blank issue rows per
                    customer; collectors add or delete rows freely.

  SPRForecast_<CC>  Month | Country | Customer | Payments In House Pending |
                    Expected Payments | Forecasted Overdue EOM |
                    Forecasted GT60 EOM | Forecasted GT90 EOM
                    One row per customer.

Table names carry the country suffix because Excel requires table names to be
unique across a workbook. The report script finds them by the "SPRIssues_" /
"SPRForecast_" prefix, so adding a new country later needs no code change.

Usage:  python3 docs/build_spr_sheets.py <pairs.json> <month> <out.xlsx>
        pairs.json = [{"country": "US", "customers": ["Amazon", ...]}, ...]
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

ISSUE_ROWS_PER_CUSTOMER = 3   # starting blank issue rows; collectors add more

ISSUE_COLS = ["Month", "Country", "Customer", "Issue", "Action"]
FORECAST_COLS = ["Month", "Country", "Customer",
                 "Payments In House Pending", "Expected Payments",
                 "Forecasted Overdue EOM", "Forecasted GT60 EOM",
                 "Forecasted GT90 EOM"]

ISSUE_WIDTHS = [12, 10, 24, 60, 60]
FORECAST_WIDTHS = [12, 10, 24, 26, 20, 24, 22, 22]

HEADER_FILL = PatternFill("solid", fgColor="D40511")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
MONEY_FMT = '#,##0'


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 28


def add_table(ws, name, first_row, last_row, ncols, style="TableStyleLight9"):
    ref = f"A{first_row}:{get_column_letter(ncols)}{last_row}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True)
    ws.add_table(t)


def build(pairs, month, out_path):
    wb = Workbook()
    wb.remove(wb.active)

    readme = wb.create_sheet("READ ME")
    notes = [
        ("How to use these SPR sheets", True),
        ("", False),
        ("There is one sheet per country. Each sheet has two tables.", False),
        ("", False),
        ("ISSUES table (top): one row per issue.", True),
        ("Type a short issue in the Issue column and the matching update in Action.", False),
        ("Add a row for each additional issue, or delete rows you do not need.", False),
        ("Rows left blank are simply ignored by the report.", False),
        ("", False),
        ("FORECAST table (below): one row per customer.", True),
        ("Enter the cash forecasting figures in euros. Leave blank or 0 if not applicable.", False),
        ("", False),
        ("Month format is YYYY-MM, for example 2026-07.", True),
        ("Keep the Month, Country, and Customer values as they are so the report can match them.", False),
        ("To start a new month, copy the existing rows, paste them below, and change the Month.", False),
        ("History is kept, so past months stay in the table.", False),
    ]
    for i, (txt, bold) in enumerate(notes, start=1):
        c = readme.cell(row=i, column=1, value=txt)
        c.font = Font(bold=bold, size=12 if bold and i == 1 else 11)
    readme.column_dimensions["A"].width = 100

    for entry in pairs:
        cc = entry["country"]
        customers = entry["customers"]
        ws = wb.create_sheet(f"SPR_{cc}")

        # ---- issues table ----
        ws.append(ISSUE_COLS)
        style_header(ws, 1, len(ISSUE_COLS))
        r = 2
        for cust in customers:
            for _ in range(ISSUE_ROWS_PER_CUSTOMER):
                ws.cell(row=r, column=1, value=month)
                ws.cell(row=r, column=2, value=cc)
                ws.cell(row=r, column=3, value=cust)
                for col in (4, 5):
                    ws.cell(row=r, column=col).alignment = Alignment(
                        wrap_text=True, vertical="top")
                ws.row_dimensions[r].height = 30
                r += 1
        issues_last = r - 1
        add_table(ws, f"SPRIssues_{cc}", 1, issues_last, len(ISSUE_COLS))

        # ---- forecast table, two blank rows below ----
        fstart = issues_last + 3
        for j, h in enumerate(FORECAST_COLS, start=1):
            ws.cell(row=fstart, column=j, value=h)
        style_header(ws, fstart, len(FORECAST_COLS))
        r = fstart + 1
        for cust in customers:
            ws.cell(row=r, column=1, value=month)
            ws.cell(row=r, column=2, value=cc)
            ws.cell(row=r, column=3, value=cust)
            for col in range(4, len(FORECAST_COLS) + 1):
                ws.cell(row=r, column=col).number_format = MONEY_FMT
            r += 1
        add_table(ws, f"SPRForecast_{cc}", fstart, r - 1, len(FORECAST_COLS),
                  style="TableStyleLight11")

        for j, w in enumerate(ISSUE_WIDTHS, start=1):
            ws.column_dimensions[get_column_letter(j)].width = max(
                w, FORECAST_WIDTHS[j - 1] if j <= len(FORECAST_WIDTHS) else 0)
        for j in range(len(ISSUE_WIDTHS) + 1, len(FORECAST_COLS) + 1):
            ws.column_dimensions[get_column_letter(j)].width = FORECAST_WIDTHS[j - 1]

        ws.freeze_panes = "A2"

    wb.save(out_path)
    sheets = len(wb.sheetnames) - 1
    print(f"Saved {out_path}: {sheets} country sheets, month {month}")


if __name__ == "__main__":
    pairs_path = sys.argv[1]
    month = sys.argv[2]
    out_path = sys.argv[3]
    with open(pairs_path) as f:
        pairs = json.load(f)
    build(pairs, month, out_path)
